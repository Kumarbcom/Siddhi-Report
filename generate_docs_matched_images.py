import os
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Inches as DocxInches

def get_path(filename):
    if filename.endswith(".png") and "1784031509588" in filename:
        return os.path.join(r"C:\Users\DELL\.gemini\antigravity\brain\tempmediaStorage", filename)
    return os.path.join(r"C:\Users\DELL\.gemini\antigravity\brain\8c545920-27be-41b4-b8c3-1290228e41ed", filename)

# Mapping of components to their corresponding equipment image and LAPP product image
items = [
    {
        "component": "Drag Chains",
        "product": "LAPP SILVYN® CHAIN",
        "description": "LAPP's SILVYN® CHAIN systems are designed for cable protection and management in dynamic applications.",
        "eq_img": get_path("media__1784031395173.jpg"),
        "prod_img": get_path("media__1784031544010.png")
    },
    {
        "component": "Continuous Movement Cables",
        "product": "LAPP ÖLFLEX® SERVO FD",
        "description": "ÖLFLEX® FD cables are highly flexible cables specifically engineered for continuous movement within drag chains. Suitable for applications due to their high flexibility and durability in constantly moving parts.",
        "eq_img": get_path("media__1784031395183.jpg"),
        "prod_img": get_path("media__1784031533268.png")
    },
    {
        "component": "Control Cabinet and Cable Glands",
        "product": "LAPP SKINTOP® Cable Gland",
        "description": "SKINTOP® cable glands provide secure, strain-relieved, and often liquid-tight cable entry into control cabinets.",
        "eq_img": get_path("media__1784031395169.jpg"),
        "prod_img": get_path("media__1784031553573.png")
    },
    {
        "component": "Control Cabinet Wiring",
        "product": "LAPP ÖLFLEX® UNIPLUS",
        "description": "ÖLFLEX® UNIPLUS single-core cables are ideal for internal wiring of devices and control cabinets.",
        "eq_img": get_path("media__1784031509588.png"),
        "prod_img": get_path("media__1784031565190.png")
    },
    {
        "component": "Cable Bundling",
        "product": "LAPP KW Plastic Coil",
        "description": "KW plastic coils are used for easy and quick bundling of cables, providing mechanical protection and organization.",
        "eq_img": get_path("media__1784031395183.jpg"),
        "prod_img": get_path("media__1784031579269.png")
    },
    {
        "component": "Circular Connectors",
        "product": "LAPP EPIC® M12",
        "description": "EPIC® circular connectors, such as M12 connectors, are suitable for robust data, signal, and power connections in industrial applications, offering vibration protection and easy assembly.",
        "eq_img": get_path("media__1784031395169.jpg"),
        "prod_img": get_path("media__1784031866201.png")
    },
    {
        "component": "Laptop Connection and Data Cables",
        "product": "UNITRONIC® & ETHERLINE®",
        "description": "UNITRONIC® (Data & Communication Cables) and ETHERLINE® (Industrial Ethernet Cables) are used for reliable data transmission and network connectivity.",
        "eq_img": None,
        "prod_img": get_path("media__1784031524111.png")
    }
]

def create_ppt():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "LAPP Products Used in Equipment"
    subtitle.text = "Component Integration & Image Matching"
    
    # Content Slides
    bullet_slide_layout = prs.slide_layouts[1]
    for item in items:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = item["component"]
        
        # Adjust text box width to make room for images
        body_shape.width = Inches(4.5)
        
        tf = body_shape.text_frame
        
        p1 = tf.paragraphs[0]
        p1.text = f"Matching LAPP Product: {item['product']}"
        p1.font.bold = True
        
        p2 = tf.add_paragraph()
        p2.text = item["description"]
        p2.level = 1
        
        # Add Equipment Image if exists
        try:
            if item["eq_img"] and os.path.exists(item["eq_img"]):
                slide.shapes.add_picture(item["eq_img"], Inches(5), Inches(1.5), width=Inches(4.5))
        except Exception as e:
            print(f"Failed to add eq image {item['eq_img']}: {e}")
            
        # Add Product Image if exists
        try:
            if item["prod_img"] and os.path.exists(item["prod_img"]):
                # If there's an eq image, put product image below it. Otherwise put it near the top.
                top = Inches(4.5) if (item["eq_img"] and os.path.exists(item["eq_img"])) else Inches(2.0)
                slide.shapes.add_picture(item["prod_img"], Inches(5.2), top, width=Inches(4))
        except Exception as e:
            print(f"Failed to add prod image {item['prod_img']}: {e}")
            
    prs.save('LAPP_Equipment_Presentation_v3.pptx')

def create_word():
    doc = Document()
    doc.add_heading('LAPP Products Used in Equipment', 0)
    
    for i, item in enumerate(items, 1):
        doc.add_heading(f"{i}. {item['component']}", level=1)
        
        p = doc.add_paragraph()
        p.add_run("Matching LAPP Products: ").bold = True
        p.add_run(item['product'])
        
        doc.add_paragraph(item['description'])
        
        # Add Images
        if (item["eq_img"] and os.path.exists(item["eq_img"])) or (item["prod_img"] and os.path.exists(item["prod_img"])):
            table = doc.add_table(rows=1, cols=2)
            row_cells = table.rows[0].cells
            
            if item["eq_img"] and os.path.exists(item["eq_img"]):
                try:
                    p = row_cells[0].paragraphs[0]
                    run = p.add_run()
                    run.add_picture(item["eq_img"], width=DocxInches(3.0))
                    p.add_run("\nEquipment")
                except:
                    pass
            
            if item["prod_img"] and os.path.exists(item["prod_img"]):
                try:
                    p = row_cells[1].paragraphs[0]
                    run = p.add_run()
                    run.add_picture(item["prod_img"], width=DocxInches(3.0))
                    p.add_run("\nLAPP Product")
                except:
                    pass
        
        doc.add_paragraph() # Spacing
        
    doc.save('LAPP_Equipment_Report_v3.docx')

if __name__ == "__main__":
    create_ppt()
    create_word()
    print("Successfully generated PPTX and DOCX files with matched images per category.")
