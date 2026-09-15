import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Inches as DocxInches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_path(filename):
    if filename.endswith(".png") and "1784031509588" in filename:
        return os.path.join(r"C:\Users\DELL\.gemini\antigravity\brain\tempmediaStorage", filename)
    return os.path.join(r"C:\Users\DELL\.gemini\antigravity\brain\8c545920-27be-41b4-b8c3-1290228e41ed", filename)

items = [
    {
        "component": "Drag Chains",
        "product": "LAPP SILVYN® CHAIN",
        "description": "LAPP's SILVYN® CHAIN systems are designed for cable protection and management in dynamic applications.",
        "eq_img": get_path("media__1784031395173.jpg"),
        "prod_img": get_path("media__1784031544010.png")
    },
    {
        "component": "Continuous Movement Cables",
        "product": "LAPP ÖLFLEX® SERVO FD",
        "description": "ÖLFLEX® FD cables are highly flexible cables specifically engineered for continuous movement within drag chains. Suitable for applications due to their high flexibility and durability in constantly moving parts.",
        "eq_img": get_path("media__1784031395183.jpg"),
        "prod_img": get_path("media__1784031533268.png")
    },
    {
        "component": "Control Cabinet and Cable Glands",
        "product": "LAPP SKINTOP® Cable Gland",
        "description": "SKINTOP® cable glands provide secure, strain-relieved, and often liquid-tight cable entry into control cabinets.",
        "eq_img": get_path("media__1784031395169.jpg"),
        "prod_img": get_path("media__1784031553573.png")
    },
    {
        "component": "Control Cabinet Wiring",
        "product": "LAPP ÖLFLEX® UNIPLUS",
        "description": "ÖLFLEX® UNIPLUS single-core cables are ideal for internal wiring of devices and control cabinets.",
        "eq_img": get_path("media__1784031509588.png"),
        "prod_img": get_path("media__1784031565190.png")
    },
    {
        "component": "Cable Bundling",
        "product": "LAPP KW Plastic Coil",
        "description": "KW plastic coils are used for easy and quick bundling of cables, providing mechanical protection and organization.",
        "eq_img": get_path("media__1784031395183.jpg"),
        "prod_img": get_path("media__1784031579269.png")
    },
    {
        "component": "Circular Connectors",
        "product": "LAPP EPIC® M12",
        "description": "EPIC® circular connectors, such as M12 connectors, are suitable for robust data, signal, and power connections in industrial applications, offering vibration protection and easy assembly.",
        "eq_img": get_path("media__1784031395169.jpg"),
        "prod_img": get_path("media__1784031866201.png")
    },
    {
        "component": "Laptop Connection and Data Cables",
        "product": "UNITRONIC® & ETHERLINE®",
        "description": "UNITRONIC® (Data & Communication Cables) and ETHERLINE® (Industrial Ethernet Cables) are used for reliable data transmission and network connectivity.",
        "eq_img": None,
        "prod_img": get_path("media__1784031524111.png")
    }
]

def create_ppt():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "LAPP Products Used in Equipment"
    subtitle.text = "Component Integration & Image Matching"
    
    # Content Slides
    blank_layout = prs.slide_layouts[5] # Title only
    
    for item in items:
        slide = prs.slides.add_slide(blank_layout)
        title_shape = slide.shapes.title
        title_shape.text = item["component"]
        
        # Add text box for description on the left
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(4.5)
        height = Inches(4.0)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p = tf.add_paragraph()
        p.text = f"Matching LAPP Product: {item['product']}"
        p.font.bold = True
        p.font.size = Pt(20)
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{item['description']}"
        p2.font.size = Pt(18)
        
        # Add images on the right
        right_col_left = Inches(5.5)
        
        # Equipment Image
        if item["eq_img"] and os.path.exists(item["eq_img"]):
            # Label
            lbl = slide.shapes.add_textbox(right_col_left, Inches(1.3), Inches(4.0), Inches(0.5))
            lbl.text_frame.text = "Application in Equipment:"
            lbl.text_frame.paragraphs[0].font.bold = True
            lbl.text_frame.paragraphs[0].font.size = Pt(14)
            # Image
            try:
                # Constrain by height so tall images fit
                slide.shapes.add_picture(item["eq_img"], right_col_left, Inches(1.7), height=Inches(2.4))
            except Exception as e:
                print(f"Failed eq img: {e}")
                
        # LAPP Product Image
        if item["prod_img"] and os.path.exists(item["prod_img"]):
            prod_top = Inches(4.5) if (item["eq_img"] and os.path.exists(item["eq_img"])) else Inches(2.0)
            
            # Label
            lbl = slide.shapes.add_textbox(right_col_left, prod_top - Inches(0.4), Inches(4.0), Inches(0.5))
            lbl.text_frame.text = "LAPP Product:"
            lbl.text_frame.paragraphs[0].font.bold = True
            lbl.text_frame.paragraphs[0].font.size = Pt(14)
            # Image
            try:
                slide.shapes.add_picture(item["prod_img"], right_col_left, prod_top, height=Inches(2.4))
            except Exception as e:
                print(f"Failed prod img: {e}")
            
    prs.save('LAPP_Equipment_Presentation_v4.pptx')

def create_word():
    doc = Document()
    doc.add_heading('LAPP Products Used in Equipment', 0)
    
    for i, item in enumerate(items, 1):
        doc.add_heading(f"{i}. {item['component']}", level=1)
        
        p = doc.add_paragraph()
        p.add_run("Matching LAPP Products: ").bold = True
        p.add_run(item['product'])
        
        doc.add_paragraph(item['description'])
        
        # Add Images side by side using a table
        if (item["eq_img"] and os.path.exists(item["eq_img"])) or (item["prod_img"] and os.path.exists(item["prod_img"])):
            table = doc.add_table(rows=2, cols=2)
            
            # Headers
            if item["eq_img"] and os.path.exists(item["eq_img"]):
                table.rows[0].cells[0].text = "Application in Equipment:"
            if item["prod_img"] and os.path.exists(item["prod_img"]):
                table.rows[0].cells[1].text = "LAPP Product:"
                
            # Images
            if item["eq_img"] and os.path.exists(item["eq_img"]):
                try:
                    p = table.rows[1].cells[0].paragraphs[0]
                    run = p.add_run()
                    # Use a fixed width for word so it scales nicely side by side
                    run.add_picture(item["eq_img"], width=DocxInches(2.5))
                except:
                    pass
            
            if item["prod_img"] and os.path.exists(item["prod_img"]):
                try:
                    p = table.rows[1].cells[1].paragraphs[0]
                    run = p.add_run()
                    run.add_picture(item["prod_img"], width=DocxInches(2.5))
                except:
                    pass
        
        doc.add_paragraph() # Spacing
        doc.add_page_break()
        
    doc.save('LAPP_Equipment_Report_v4.docx')

if __name__ == "__main__":
    create_ppt()
    create_word()
    print("Successfully generated v4 documents.")
