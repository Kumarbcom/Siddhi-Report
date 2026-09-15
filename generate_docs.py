import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Pt as DocxPt

# Data for the presentation and report
items = [
    {
        "component": "Drag Chains",
        "product": "LAPP SILVYN® CHAIN",
        "description": "LAPP's SILVYN® CHAIN systems are designed for cable protection and management in dynamic applications."
    },
    {
        "component": "Continuous Movement Cables",
        "product": "LAPP ÖLFLEX® SERVO FD",
        "description": "ÖLFLEX® FD cables are highly flexible cables specifically engineered for continuous movement within drag chains. Suitable for applications due to their high flexibility and durability in constantly moving parts."
    },
    {
        "component": "Control Cabinet and Cable Glands",
        "product": "LAPP SKINTOP® Cable Gland",
        "description": "SKINTOP® cable glands provide secure, strain-relieved, and often liquid-tight cable entry into control cabinets."
    },
    {
        "component": "Control Cabinet Wiring",
        "product": "LAPP ÖLFLEX® UNIPLUS",
        "description": "ÖLFLEX® UNIPLUS single-core cables are ideal for internal wiring of devices and control cabinets."
    },
    {
        "component": "Cable Bundling",
        "product": "LAPP KW Plastic Coil",
        "description": "KW plastic coils are used for easy and quick bundling of cables, providing mechanical protection and organization."
    },
    {
        "component": "Circular Connectors",
        "product": "LAPP EPIC® M12",
        "description": "EPIC® circular connectors, such as M12 connectors, are suitable for robust data, signal, and power connections in industrial applications, offering vibration protection and easy assembly."
    },
    {
        "component": "Laptop Connection and Data Cables",
        "product": "UNITRONIC® & ETHERLINE®",
        "description": "UNITRONIC® (Data & Communication Cables) and ETHERLINE® (Industrial Ethernet Cables) are used for reliable data transmission and network connectivity."
    }
]

def create_ppt():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "LAPP Products Used in Equipment"
    subtitle.text = "Component Integration Overview"
    
    # Content Slides
    bullet_slide_layout = prs.slide_layouts[1]
    for item in items:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = item["component"]
        tf = body_shape.text_frame
        
        p1 = tf.paragraphs[0]
        p1.text = f"Matching LAPP Product: {item['product']}"
        p1.font.bold = True
        
        p2 = tf.add_paragraph()
        p2.text = item["description"]
        p2.level = 1
        
    prs.save('LAPP_Equipment_Presentation.pptx')

def create_word():
    doc = Document()
    doc.add_heading('LAPP Products Used in Equipment', 0)
    
    for i, item in enumerate(items, 1):
        doc.add_heading(f"{i}. {item['component']}", level=1)
        
        p = doc.add_paragraph()
        p.add_run("Matching LAPP Products: ").bold = True
        p.add_run(item['product'])
        
        doc.add_paragraph(item['description'])
        doc.add_paragraph() # Add some spacing
        
    doc.save('LAPP_Equipment_Report.docx')

if __name__ == "__main__":
    create_ppt()
    create_word()
    print("Successfully generated PPTX and DOCX files.")
